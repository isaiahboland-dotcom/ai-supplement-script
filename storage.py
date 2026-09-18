"""
Output locations:
  output/research/<name>.json   full record (raw research, review, sources, final fields)
  Grist                         if GRIST_URL/DOC_ID/API_KEY are set: add-or-update by Supplement_Name
  output/supplements_research.xlsx   otherwise, or when Grist fails (same layout as example.xlsx)
"""

import json
import re
from pathlib import Path
from typing import Dict

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

import config

COLUMNS = [
    "Supplement_Name",
    "Safe_Effective_Dose_Range",
    "Safe_Use_Duration_Schedule",
    "Benefits",
    "Negative_Effects",
    "Advisories",
    "Recommendations_for_Effectiveness",
    "Sources",
]

# widths copied from example.xlsx
COLUMN_WIDTHS = [39, 39, 39, 39, 39, 39, 39, 39]


def sanitize_filename(name: str) -> str:
    cleaned = re.sub(r"[^\w\-]+", "_", name.strip())
    return cleaned[:80] or "supplement"


def record_path(supplement: str) -> Path:
    return config.RESEARCH_DIR / f"{sanitize_filename(supplement).lower()}.json"


def is_done(supplement: str) -> bool:
    path = record_path(supplement)
    if not path.exists():
        return False
    try:
        return json.loads(path.read_text(encoding="utf-8")).get("status") == "complete"
    except (OSError, json.JSONDecodeError):
        return False


def save_record(supplement: str, record: dict):
    config.RESEARCH_DIR.mkdir(parents=True, exist_ok=True)
    record_path(supplement).write_text(
        json.dumps(record, indent=2, ensure_ascii=False), encoding="utf-8"
    )


def _new_workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Supplements"
    ws.append(COLUMNS)

    for i, width in enumerate(COLUMN_WIDTHS, 1):
        letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[letter].width = width
        ws.cell(row=1, column=i).font = Font(bold=True)

    ws.freeze_panes = "A2"
    return wb


def upsert_excel(fields: Dict[str, str]):
    config.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    path = config.EXCEL_FILE

    wb = load_workbook(path) if path.exists() else _new_workbook()
    ws = wb.active

    header = [c.value for c in ws[1]]
    missing = [c for c in COLUMNS if c not in header]
    if missing:
        raise RuntimeError(
            f"{path} is missing columns {missing}; move it aside so a fresh one is created."
        )
    col_index = {name: header.index(name) + 1 for name in COLUMNS}

    target = fields["Supplement_Name"].strip().lower()
    row_num = None
    for r in range(2, ws.max_row + 1):
        val = ws.cell(row=r, column=col_index["Supplement_Name"]).value
        if val and str(val).strip().lower() == target:
            row_num = r
            break

    # nothing matched, so append. back up over trailing blank rows first
    if row_num is None:
        row_num = ws.max_row + 1
        while row_num > 2 and not any(c.value for c in ws[row_num - 1]):
            row_num -= 1

    wrap = Alignment(wrap_text=True, vertical="top")
    for name in COLUMNS:
        cell = ws.cell(row=row_num, column=col_index[name], value=fields.get(name, ""))
        cell.alignment = wrap

    try:
        wb.save(path)
    except PermissionError as e:
        raise RuntimeError(
            f"Could not write {path} (is it open in a spreadsheet app?)"
        ) from e


def upsert_grist(fields: Dict[str, str]):
    # PUT with a require clause so Grist does the add-or-update itself
    url = (
        f"{config.GRIST_URL}/api/docs/{config.GRIST_DOC_ID}"
        f"/tables/{config.GRIST_TABLE_ID}/records"
    )

    payload = {
        "records": [
            {
                "require": {"Supplement_Name": fields["Supplement_Name"]},
                "fields": {k: v for k, v in fields.items() if k != "Supplement_Name"},
            }
        ]
    }

    resp = requests.put(
        url,
        headers={
            "Authorization": f"Bearer {config.GRIST_API_KEY}",
            "Content-Type": "application/json",
        },
        json=payload,
        timeout=60,
    )

    if resp.status_code >= 400:
        raise RuntimeError(f"Grist HTTP {resp.status_code}: {resp.text[:300]}")
